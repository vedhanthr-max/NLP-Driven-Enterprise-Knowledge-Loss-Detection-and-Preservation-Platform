import os
import hashlib
from datetime import datetime, timezone
from typing import Tuple, List, Optional
from sqlalchemy.orm import Session
from fastapi import UploadFile

from app.core.config import settings
from app.core.exceptions import (
    AppException, DocumentNotFoundException, InvalidFileFormatException, FileTooLargeException
)
from app.models.entities import (
    Document, DocumentChunk, KnowledgeItem, Entity, Keyword, Embedding, User
)
from app.nlp.preprocessor import clean_text, chunk_text
from app.nlp.ner import extract_entities
from app.nlp.keywords import extract_keywords
from app.nlp.topic_extractor import extract_topics
from app.nlp.embeddings import generate_embedding
from app.risk.scoring import estimate_documentation_gap, estimate_uniqueness_from_experts
from app.services.risk_service import RiskService

class DocumentService:
    @staticmethod
    def extract_text_from_file(file_path: str, file_type: str) -> Tuple[str, int]:
        """Extract plain text and page count from PDF, DOCX, or TXT."""
        text = ""
        page_count = 1

        if file_type == "pdf":
            try:
                import fitz  # PyMuPDF
                doc = fitz.open(file_path)
                page_count = len(doc)
                extracted_pages = []
                for page in doc:
                    extracted_pages.append(page.get_text())
                text = "\n\n".join(extracted_pages)
                doc.close()
            except Exception as e:
                raise AppException(400, "ERR-DOC-004", f"Corrupted or unreadable PDF: {str(e)}")

        elif file_type == "docx":
            try:
                import docx
                doc = docx.Document(file_path)
                paragraphs = [p.text for p in doc.paragraphs if p.text]
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                        if row_text:
                            paragraphs.append(row_text)
                text = "\n\n".join(paragraphs)
                page_count = max(1, len(paragraphs) // 10)
            except Exception as e:
                raise AppException(400, "ERR-DOC-004", f"Corrupted or unreadable DOCX: {str(e)}")

        elif file_type in ("txt", "md"):
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
                page_count = max(1, len(text.splitlines()) // 40)
            except Exception as e:
                raise AppException(400, "ERR-DOC-004", f"Text file extraction failed: {str(e)}")
        else:
            raise InvalidFileFormatException(file_type, settings.ALLOWED_EXTENSIONS)

        return clean_text(text), page_count

    @classmethod
    async def process_document_upload(
        cls,
        db: Session,
        file: UploadFile,
        title: str,
        description: Optional[str],
        current_user: User
    ) -> Document:
        """Handle file upload, validation, extraction, chunking, NLP, and knowledge creation."""
        # 1. Validation
        ext = file.filename.split(".")[-1].lower() if "." in file.filename else ""
        if ext not in settings.ALLOWED_EXTENSIONS:
            raise InvalidFileFormatException(ext, settings.ALLOWED_EXTENSIONS)

        contents = await file.read()
        file_size = len(contents)
        if file_size == 0:
            raise AppException(400, "ERR-DOC-005", "Uploaded file is empty (0 bytes)")
        if file_size > settings.MAX_FILE_SIZE:
            raise FileTooLargeException(file_size, settings.MAX_FILE_SIZE)

        # 2. Deduplication Hash
        content_hash = hashlib.sha256(contents).hexdigest()
        existing_doc = db.query(Document).filter(
            Document.organization_id == current_user.organization_id,
            Document.content_hash == content_hash,
            Document.deleted_at.is_(None)
        ).first()

        if existing_doc:
            return existing_doc

        # 3. Save File to disk
        os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
        safe_filename = f"{current_user.organization_id}_{int(datetime.now(timezone.utc).timestamp())}_{file.filename}"
        file_path = os.path.join(settings.UPLOAD_DIR, safe_filename)
        with open(file_path, "wb") as f:
            f.write(contents)

        # 4. Extract Text
        raw_text, pages = cls.extract_text_from_file(file_path, ext)

        # 5. Create Document Record
        doc = Document(
            organization_id=current_user.organization_id,
            title=title or file.filename,
            description=description,
            file_name=file.filename,
            file_path=file_path,
            file_size=file_size,
            file_type=ext,
            mime_type=file.content_type,
            content_hash=content_hash,
            pages=pages,
            uploaded_by=current_user.user_id,
            extraction_status="completed",
            processed_at=datetime.now(timezone.utc)
        )
        db.add(doc)
        db.flush()

        # 6. Chunking & Embeddings
        chunks_data = chunk_text(raw_text, chunk_size=settings.CHUNK_SIZE, overlap=settings.CHUNK_OVERLAP)
        created_chunks = []
        for c in chunks_data:
            chunk = DocumentChunk(
                document_id=doc.document_id,
                chunk_index=c["chunk_index"],
                chunk_text=c["chunk_text"],
                token_count=c["token_count"],
                start_position=c["start_position"],
                end_position=c["end_position"],
                page_number=c["page_number"]
            )
            db.add(chunk)
            created_chunks.append(chunk)
        db.flush()

        # 7. NLP Extraction
        entities_data = extract_entities(raw_text)
        keywords_data = extract_keywords(raw_text, top_n=15)
        primary_topic, _ = extract_topics(raw_text)

        # 8. Create Associated Knowledge Item
        doc_snippet = raw_text[:350].strip() if raw_text else "No content"
        knowledge_item = KnowledgeItem(
            organization_id=current_user.organization_id,
            title=title or file.filename,
            description=description or f"Extracted knowledge from {file.filename}: {doc_snippet}...",
            category=primary_topic,
            tags=[kw["keyword"] for kw in keywords_data[:5]],
            source_document_id=doc.document_id,
            knowledge_type="procedure",
            status="active",
            confidence_score=0.92,
            extraction_method="nlp_extracted",
            created_by=current_user.user_id
        )
        db.add(knowledge_item)
        db.flush()

        # 9. Store Entities & Keywords
        for ent in entities_data:
            db.add(Entity(
                knowledge_id=knowledge_item.knowledge_id,
                document_id=doc.document_id,
                entity_name=ent["entity_name"],
                entity_type=ent["entity_type"],
                confidence=ent["confidence"],
                context=ent["context"]
            ))

        for kw in keywords_data:
            db.add(Keyword(
                knowledge_id=knowledge_item.knowledge_id,
                keyword=kw["keyword"],
                frequency=kw["frequency"],
                tf_idf_score=kw["tf_idf_score"]
            ))

        # 10. Store Chunk Embeddings
        for chunk in created_chunks:
            vec = generate_embedding(chunk.chunk_text)
            embedding_record = Embedding(
                knowledge_id=knowledge_item.knowledge_id,
                document_chunk_id=chunk.chunk_id,
                embedding_model=settings.EMBEDDING_MODEL
            )
            embedding_record.set_vector(vec)
            db.add(embedding_record)

        # 11. Calculate Initial Automated Risk Assessment
        doc_gap = estimate_documentation_gap(raw_text)
        person_count = sum(1 for e in entities_data if e["entity_type"] == "PERSON")
        uniqueness = estimate_uniqueness_from_experts(max(1, person_count))

        # Baseline automatic scores from document context
        crit_score = 75.0 if any(k in raw_text.lower() for k in ["payment", "production", "critical", "legacy", "database"]) else 50.0
        dep_score = 70.0 if any(k in raw_text.lower() for k in ["core", "enterprise", "department", "api"]) else 45.0
        avail_score = 65.0 if uniqueness > 70 else 30.0

        RiskService.calculate_and_save_risk(
            db=db,
            knowledge_id=knowledge_item.knowledge_id,
            criticality=crit_score,
            uniqueness=uniqueness,
            dependency=dep_score,
            documentation_gap=doc_gap,
            availability_risk=avail_score,
            assessment_method="rule_based",
            current_user_id=current_user.user_id
        )

        db.commit()
        db.refresh(doc)
        return doc
